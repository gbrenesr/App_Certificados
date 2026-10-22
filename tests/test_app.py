# tests/test_app.py — pruebas automáticas de la app de certificados.
# Se usa el cliente de pruebas de Flask: simula un navegador sin abrir la app de verdad.

import os
from io import BytesIO
from pathlib import Path
import re

import openpyxl
import pytest

# Ladrillo: variables de entorno de prueba ANTES de importar la app.
# Así las pruebas no dependen del archivo .env de la máquina.
os.environ["ADMIN_CORREO"] = "admin@prueba.cr"
os.environ["ADMIN_CLAVE"] = "clave-admin-prueba"
os.environ["USUARIO_CORREO"] = "usuario@prueba.cr"
os.environ["USUARIO_CLAVE"] = "clave-usuario-prueba"
os.environ["SECRET_KEY"] = "clave-secreta-solo-para-pruebas"

from app import app  # noqa: E402  (se importa después de definir el entorno)

# Ladrillo: rutas de la Bodega (los archivos originales NO se modifican)
BASE = Path(__file__).resolve().parent.parent
RUTA_MAESTRO = BASE / "Insumos" / "Maestro_Estudiantes.xlsx"
RUTA_EVALUACIONES = BASE / "Insumos" / "Registro_Evaluaciones.xlsx"


@pytest.fixture
def cliente():
    """Crea un navegador de mentira limpio para cada prueba."""
    app.config["TESTING"] = True
    return app.test_client()


def entrar(cliente, correo, clave):
    """Ladrillo: función auxiliar que simula escribir el correo y la clave."""
    return cliente.post("/login", data={"correo": correo, "clave": clave})


def contadores_de(html):
    """Saca los 3 números de los contadores de arriba, en orden."""
    # Ladrillo: BUCLE con expresiones regulares para leer los <strong> de cada tarjeta
    patrones = [
        r'class="tarjeta verde"><strong>(\d+)</strong>',
        r'class="tarjeta azul"><strong>(\d+)</strong>',
        r'class="tarjeta roja"><strong>(\d+)</strong>',
    ]
    return [int(re.search(p, html).group(1)) for p in patrones]


def registro_con_camila_en_40():
    """Arma en memoria una copia del Registro con las notas de Camila en 40.

    Ladrillo: BytesIO es la "memoria USB imaginaria" en RAM; el archivo original
    de Insumos solo se lee, nunca se guarda ni se modifica.
    """
    libro = openpyxl.load_workbook(RUTA_EVALUACIONES)
    hoja = libro["Evaluaciones"]
    cabeceras = [c.value for c in hoja[1]]
    columna_id = cabeceras.index("Identificacion") + 1
    columna_nota = cabeceras.index("Nota") + 1

    # Ladrillo: BUCLE por las filas buscando a Camila (identificación 101230456)
    for numero_fila in range(2, hoja.max_row + 1):
        identificacion = str(hoja.cell(row=numero_fila, column=columna_id).value).strip()
        if identificacion == "101230456":
            hoja.cell(row=numero_fila, column=columna_nota).value = 40

    memoria = BytesIO()
    libro.save(memoria)
    memoria.seek(0)  # rebobinar para poder leerla desde el inicio
    return memoria


def archivos_para_subir(registro_modificado):
    """Ladrillo: diccionario con los dos archivos tal como los manda un formulario."""
    return {
        "accion": "ver",
        "maestro": (open(RUTA_MAESTRO, "rb"), "Maestro_Estudiantes.xlsx"),
        "evaluaciones": (registro_modificado, "Registro_Evaluaciones.xlsx"),
    }


def test_sin_sesion_redirige_a_login(cliente):
    """Simula abrir la app sin haber entrado: debe mandar a la pantalla de entrada."""
    respuesta = cliente.get("/")
    assert respuesta.status_code == 302  # Ladrillo: 302 = redirección
    assert "/login" in respuesta.headers["Location"]


def test_clave_incorrecta_no_entra(cliente):
    """Simula escribir un correo y una clave que no existen: queda en el login con error."""
    respuesta = entrar(cliente, "nadie@prueba.cr", "clave-equivocada")
    assert respuesta.status_code == 200  # sigue en la pantalla de entrada
    html = respuesta.get_data(as_text=True)
    assert "Correo o contraseña incorrectos" in html


def test_usuario_no_puede_cargar(cliente):
    """Simula entrar como usuario normal e intentar usar el panel de carga: debe dar 403."""
    entrar(cliente, os.environ["USUARIO_CORREO"], os.environ["USUARIO_CLAVE"])
    respuesta = cliente.post("/cargar", data={"accion": "ver"})
    assert respuesta.status_code == 403  # Ladrillo: 403 = prohibido


def test_excel_modificado_cambia_contadores(cliente):
    """Simula que un admin sube un Registro con Camila en 40 y ve los contadores nuevos."""
    entrar(cliente, os.environ["ADMIN_CORREO"], os.environ["ADMIN_CLAVE"])
    registro = registro_con_camila_en_40()
    datos = archivos_para_subir(registro)
    try:
        respuesta = cliente.post("/cargar", data=datos, content_type="multipart/form-data")
    finally:
        # Ladrillo: cerrar el archivo del Maestro que se abrió para subirlo
        datos["maestro"][0].close()

    assert respuesta.status_code == 200
    html = respuesta.get_data(as_text=True)

    # Ladrillo: con Camila en 40 pasa de Aprobación a Participación
    assert contadores_de(html) == [18, 4, 1]
    # Ladrillo: se confirma que se usaron los archivos subidos y se avisa que no se guardan
    assert "Archivos subidos" in html
    assert "Está viendo datos subidos" in html
