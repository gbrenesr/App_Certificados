# app.py — LA COCINA (Academia Horizonte)
# Toda la lógica vive aquí. Las páginas HTML solo muestran lo ya calculado.
# Ladrillos señalados en español: variables, tipos, funciones, bucles y condicionales.

from pathlib import Path  # Ladrillo: tipo para manejar rutas de archivos
from collections import Counter  # Ladrillo: tipo contador para los contadores de arriba
from functools import wraps  # Ladrillo: herramienta para crear el decorador requiere_login
import io  # Ladrillo: io.BytesIO = un archivo que vive en la memoria, no en el disco
import os  # Ladrillo: librería para leer variables de entorno (os.environ.get)
import secrets  # Ladrillo: comparación segura de contraseñas (sin filtrar tiempos)
import openpyxl  # Ladrillo: librería para leer los Excel (Bodega)
from dotenv import load_dotenv  # Ladrillo: lee el archivo .env en la máquina local
# Ladrillo: session = memoria del usuario; redirect/url_for = navegación; Flask = servidor web
# send_file = enviar un archivo como descarga sin guardarlo en disco
from flask import Flask, render_template, request, session, redirect, url_for, send_file
from openpyxl.styles import Font, PatternFill  # Ladrillo: estilos para el Excel de descarga
from openpyxl.utils import get_column_letter  # Ladrillo: para ajustar el ancho de columnas

# Ladrillo: cargar las variables del archivo .env (si no existe, no pasa nada).
# En Vercel las variables se configuran en el panel del proyecto, no en este archivo.
load_dotenv()

# Ladrillo: variables (texto con las rutas; la Bodega NO se modifica)
BASE = Path(__file__).resolve().parent
RUTA_MAESTRO = BASE / "Insumos" / "Maestro_Estudiantes.xlsx"
RUTA_EVALS = BASE / "Insumos" / "Registro_Evaluaciones.xlsx"

# Ladrillo: creación de la aplicación Flask (tipo Flask)
app = Flask(__name__)

# Ladrillo: clave secreta para firmar la cookie de sesión.
# En producción (Vercel) SIEMPRE debe venir de la variable de entorno SECRET_KEY.
# El valor de respaldo es solo para desarrollo local.
app.secret_key = os.environ.get("SECRET_KEY", "clave-solo-para-desarrollo-local")

# Ladrillo: límite de tamaño de la petición completa (4 MB).
# Vercel rechaza peticiones de más de ~4.5 MB, así que nos quedamos por debajo.
app.config["MAX_CONTENT_LENGTH"] = 4 * 1024 * 1024

# Ladrillo: variable de texto con el origen cuando se usan los Excel de la carpeta
ORIGEN_FABRICA = "Datos de fábrica (carpeta Insumos/)"


def cargar_usuarios():
    """Lee los 2 usuarios desde variables de entorno. Si faltan, usa datos de demostración."""
    # Ladrillo: variables de entorno (texto o None si no existen)
    admin_correo = os.environ.get("ADMIN_CORREO")
    admin_clave = os.environ.get("ADMIN_CLAVE")
    usuario_correo = os.environ.get("USUARIO_CORREO")
    usuario_clave = os.environ.get("USUARIO_CLAVE")

    # Ladrillo: CONDICIONAL que detecta si falta alguna variable -> modo demostración
    modo_demo = not (admin_correo and admin_clave and usuario_correo and usuario_clave)

    # Ladrillo: diccionario con los 2 usuarios, cada uno con correo, clave y rol
    usuarios = {}
    # Ladrillo: cada usuario se arma con lo que venga del entorno o su valor de demostración
    usuarios["admin"] = {
        "correo": admin_correo or "admin@academia.cr",
        "clave": admin_clave or "admin123",
        "rol": "admin",
    }
    usuarios["usuario"] = {
        "correo": usuario_correo or "usuario@academia.cr",
        "clave": usuario_clave or "usuario123",
        "rol": "usuario",
    }
    return usuarios, modo_demo


# Ladrillo: se cargan una sola vez al iniciar la app
USUARIOS, MODO_DEMO = cargar_usuarios()


def requiere_login(funcion):
    """Decorador: si no hay sesión iniciada, manda a /login antes de mostrar datos."""
    @wraps(funcion)
    def envoltura(*args, **kwargs):
        # Ladrillo: CONDICIONAL que revisa si la sesión tiene el correo guardado
        if "correo" not in session:
            return redirect(url_for("login"))
        return funcion(*args, **kwargs)
    return envoltura


def solo_admin(funcion):
    """Decorador: solo deja pasar si el rol de la sesión es "admin"; si no, responde 403."""
    @wraps(funcion)
    def envoltura(*args, **kwargs):
        # Ladrillo: CONDICIONAL que compara el rol guardado en la sesión
        if session.get("rol") != "admin":
            # Ladrillo: respuesta con código 403 (prohibido) y página amable
            return render_template("sin_permiso.html"), 403
        return funcion(*args, **kwargs)
    return envoltura


# Ladrillo: FUNCIÓN que calcula el promedio con un BUCLE (suma / cantidad)
def calcular_promedio(valores):
    # Ladrillo: variable de tipo número, empieza en 0
    total = 0.0
    # Ladrillo: BUCLE que recorre cada valor de la lista
    for v in valores:  # v es de tipo número (float)
        total = total + float(v)  # suma acumulada
    cantidad = len(valores)  # Ladrillo: variable entera = módulos cursados
    # Ladrillo: CONDICIONAL que evita dividir entre cero
    if cantidad == 0:
        return 0.0
    return total / cantidad


# Ladrillo: FUNCIÓN que decide el certificado con CONDICIONALES (límites INCLUYEN el valor)
def decidir_certificado(promedio, asistencia):
    # Ladrillo: variables de tipo número (float)
    prom = float(promedio)
    asis = float(asistencia)
    # Ladrillo: CONDICIONALES en orden (primero el caso sin certificado)
    if asis < 80:  # si la asistencia es menor a 80, no hay certificado
        return "Sin certificado"
    if prom >= 70:  # asistencia >= 80 y promedio >= 70
        return "Aprobación"
    return "Participación"  # asistencia >= 80 pero promedio < 70


def cargar_datos(archivo_maestro=None, archivo_evaluaciones=None):
    """Lee los Excel y prepara resultados + contadores + inconsistencias. Todo el cálculo aquí.

    Si no se le pasa ningún archivo, lee los Excel de la carpeta Insumos (datos de fábrica).
    Si se le pasa un archivo (por ejemplo io.BytesIO con un Excel subido), lee ese.
    """
    # Ladrillo: si no llega archivo, se usa la ruta de la Bodega; si llega, se usa ese archivo
    fuente_maestro = archivo_maestro if archivo_maestro is not None else RUTA_MAESTRO
    fuente_evals = archivo_evaluaciones if archivo_evaluaciones is not None else RUTA_EVALS

    # --- Leer Maestro: diccionario {identificacion (texto): datos} ---
    wb_m = openpyxl.load_workbook(fuente_maestro, data_only=True)
    ws_m = wb_m["Estudiantes"]
    cab_m = [c.value for c in ws_m[1]]
    im = {nombre: i for i, nombre in enumerate(cab_m)}  # Ladrillo: diccionario columna->posición
    maestro = {}  # Ladrillo: variable de tipo diccionario
    # Ladrillo: BUCLE por las filas del Maestro (desde la fila 2, la 1 es el encabezado)
    for fila in ws_m.iter_rows(min_row=2, values_only=True):
        ident = str(fila[im["Identificacion"]]).strip()  # Ladrillo: texto, nunca número
        maestro[ident] = {
            "nombre": fila[im["Nombre_Completo"]],
            "programa": fila[im["Programa"]],
        }

    # --- Leer Evaluaciones: agrupar por (identificacion + programa) ---
    wb_e = openpyxl.load_workbook(fuente_evals, data_only=True)
    ws_e = wb_e["Evaluaciones"]
    cab_e = [c.value for c in ws_e[1]]
    ie = {nombre: i for i, nombre in enumerate(cab_e)}
    grupos = {}  # Ladrillo: diccionario {(id, programa): lista de filas}
    # Ladrillo: BUCLE por las filas de Evaluaciones
    for fila in ws_e.iter_rows(min_row=2, values_only=True):
        ident = str(fila[ie["Identificacion"]]).strip()
        prog = fila[ie["Programa"]]
        clave = (ident, prog)  # Ladrillo: tupla = un certificado por estudiante y programa
        # Ladrillo: CONDICIONAL para crear la lista la primera vez
        if clave not in grupos:
            grupos[clave] = []
        grupos[clave].append(fila)

    # --- Cruce: un resultado por estudiante del Maestro ---
    resultados = []  # Ladrillo: lista de diccionarios (tipo lista)
    # Ladrillo: BUCLE por cada estudiante del Maestro
    for ident, datos in maestro.items():
        clave = (ident, datos["programa"])
        evs = grupos.get(clave, [])
        # Ladrillo: CONDICIONAL sin evaluaciones
        if len(evs) == 0:
            resultados.append({
                "identificacion": ident,
                "nombre": datos["nombre"],
                "programa": datos["programa"],
                "modulos": 0,
                "promedio": 0.0,
                "asistencia": 0.0,
                "certificado": "Sin registros",
            })
        else:
            notas = [float(f[ie["Nota"]]) for f in evs]  # Ladrillo: lista de números
            asists = [float(f[ie["Asistencia_Pct"]]) for f in evs]
            prom = round(calcular_promedio(notas), 2)  # Ladrillo: uso de la FUNCIÓN
            asis = round(calcular_promedio(asists), 2)
            tipo = decidir_certificado(prom, asis)  # Ladrillo: uso de la FUNCIÓN
            resultados.append({
                "identificacion": ident,
                "nombre": datos["nombre"],
                "programa": datos["programa"],
                "modulos": len(evs),
                "promedio": prom,
                "asistencia": asis,
                "certificado": tipo,
            })
    # Ladrillo: ordenar por programa y nombre (no calcula, solo ordena)
    resultados.sort(key=lambda r: (r["programa"], r["nombre"]))

    # --- Contadores para la parte de arriba (solo certificados reales) ---
    conteo = Counter(r["certificado"] for r in resultados)  # Ladrillo: contador por tipo
    contadores = {
        "aprobacion": conteo.get("Aprobación", 0),
        "participacion": conteo.get("Participación", 0),
        "sin_certificado": conteo.get("Sin certificado", 0),
    }

    # --- Inconsistencias: las 3 detectadas ---
    inconsistencias = []
    # 1) Del Maestro SIN ninguna evaluación
    # Ladrillo: BUCLE + CONDICIONAL buscando al sin registros
    for ident, datos in maestro.items():
        if (ident, datos["programa"]) not in grupos:
            inconsistencias.append({
                "identificacion": ident,
                "nombre": datos["nombre"],
                "programa": datos["programa"],
                "detalle": "Está en el Maestro pero no tiene ninguna evaluación.",
            })
    # 2) En Evaluaciones pero NO en el Maestro (huérfano: el Maestro manda, queda en revisar)
    for (ident, prog), evs in sorted(grupos.items()):
        # Ladrillo: CONDICIONAL huérfano
        if ident not in maestro:
            notas = [float(f[ie["Nota"]]) for f in evs]
            asists = [float(f[ie["Asistencia_Pct"]]) for f in evs]
            inconsistencias.append({
                "identificacion": ident,
                "nombre": "— (no está en el Maestro)",
                "programa": prog,
                "detalle": f"Tiene {len(evs)} evaluaciones (prom {round(calcular_promedio(notas), 2)}, "
                           f"asis {round(calcular_promedio(asists), 2)}), pero no existe en el Maestro. Queda en revisar.",
            })
    # 3) Sin certificado por asistencia (caso con ceros: posible abandono)
    for r in resultados:
        # Ladrillo: CONDICIONAL para el caso de asistencia menor a 80
        if r["certificado"] == "Sin certificado":
            inconsistencias.append({
                "identificacion": r["identificacion"],
                "nombre": r["nombre"],
                "programa": r["programa"],
                "detalle": f"Promedio {r['promedio']} y asistencia {r['asistencia']}. Sin certificado por asistencia < 80.",
            })

    programas = sorted({r["programa"] for r in resultados})  # Ladrillo: lista para el filtro
    return resultados, contadores, inconsistencias, programas


def pintar_tablero(resultados, contadores, inconsistencias, programas, filtro="Todos",
                   origen=ORIGEN_FABRICA, aviso_subida=False, error=None):
    """Prepara la página del tablero. Aquí se decide qué datos recibe cada rol."""
    # Ladrillo: CONDICIONAL por rol, el usuario normal nunca recibe las inconsistencias
    if session.get("rol") != "admin":
        inconsistencias = []
    return render_template("index.html", resultados=resultados, contadores=contadores,
                           inconsistencias=inconsistencias, programas=programas, filtro=filtro,
                           origen=origen, aviso_subida=aviso_subida, error=error)


def construir_excel(resultados):
    """Arma un Excel con los resultados dentro de la memoria (io.BytesIO), sin tocar el disco."""
    libro = openpyxl.Workbook()  # Ladrillo: libro nuevo de Excel
    hoja = libro.active
    hoja.title = "Resultados"
    # Ladrillo: lista con los títulos de las columnas
    encabezados = ["Identificación", "Nombre", "Programa", "Módulos", "Promedio",
                   "Asistencia", "Certificado"]
    hoja.append(encabezados)
    # Ladrillo: BUCLE que escribe una fila por estudiante
    for r in resultados:
        hoja.append([r["identificacion"], r["nombre"], r["programa"], r["modulos"],
                     r["promedio"], r["asistencia"], r["certificado"]])
    # Ladrillo: pintar el encabezado en azul marino con letra blanca
    relleno = PatternFill("solid", fgColor="0A2A5A")
    for celda in hoja[1]:
        celda.fill = relleno
        celda.font = Font(color="FFFFFF", bold=True)
    hoja.freeze_panes = "A2"  # Ladrillo: deja fija la fila de títulos
    # Ladrillo: ancho de cada columna para que se lea bien
    anchos = [16, 28, 30, 10, 12, 12, 18]
    for numero, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(numero)].width = ancho
    # Ladrillo: BytesIO es como una memoria USB imaginaria dentro de la RAM:
    # se guarda y se lee igual que un archivo, pero nunca llega al disco y desaparece al cerrar.
    memoria = io.BytesIO()
    libro.save(memoria)  # se guarda el Excel dentro de esa memoria
    memoria.seek(0)  # rebobinar al inicio para poder enviarlo
    return memoria


# Ladrillo: ruta de entrada (formulario de correo y contraseña)
@app.route("/login", methods=["GET", "POST"])
def login():
    error = None  # Ladrillo: variable texto para el mensaje rojo de error
    # Ladrillo: CONDICIONAL si la persona ya envió el formulario
    if request.method == "POST":
        correo = request.form.get("correo", "").strip()  # Ladrillo: texto escrito
        clave = request.form.get("clave", "")
        # Ladrillo: BUCLE por los usuarios buscando el correo
        for datos in USUARIOS.values():
            # Ladrillo: CONDICIONAL que compara correo y contraseña
            if datos["correo"].lower() == correo.lower() and secrets.compare_digest(datos["clave"], clave):
                session.clear()  # limpia cualquier sesión anterior
                session["correo"] = datos["correo"]  # Ladrillo: guarda el correo en la sesión
                session["rol"] = datos["rol"]  # Ladrillo: guarda el rol en la sesión
                return redirect(url_for("index"))  # entra a la página principal
        error = "Correo o contraseña incorrectos."
    return render_template("login.html", error=error, modo_demo=MODO_DEMO)


# Ladrillo: ruta de salida (borra la sesión y vuelve a la entrada)
@app.route("/salir")
def salir():
    session.clear()  # Ladrillo: borra todo lo guardado en la sesión
    return redirect(url_for("login"))


# Ladrillo: ruta principal (página en http://localhost:5000), protegida con el portero
@app.route("/")
@requiere_login
def index():
    resultados, contadores, inconsistencias, programas = cargar_datos()
    # Ladrillo: variable texto con el filtro que eligió la persona (o "Todos")
    filtro = request.args.get("programa", "Todos")
    # Ladrillo: CONDICIONAL para filtrar en el backend (el HTML no decide)
    if filtro in programas:
        visibles = [r for r in resultados if r["programa"] == filtro]
    else:
        visibles = resultados
        filtro = "Todos"
    # Ladrillo: CONDICIONAL por rol, el usuario normal no recibe las inconsistencias.
    # Esconderlas en el HTML no basta: si viajaran, se verían con "ver código fuente".
    # El HTML solo muestra lo que aquí ya se calculó.
    return pintar_tablero(visibles, contadores, inconsistencias, programas, filtro=filtro)


# Ladrillo: ruta solo para administradores (el portero solo_admin devuelve 403 si no lo es)
@app.route("/inconsistencias")
@requiere_login
@solo_admin
def inconsistencias_admin():
    # Lleva a la sección de inconsistencias del tablero (solo el admin puede pasar)
    return redirect(url_for("index") + "#inconsistencias")


# Ladrillo: panel de carga de Excel (solo admin). Los archivos se leen EN MEMORIA.
@app.route("/cargar", methods=["POST"])
@requiere_login
@solo_admin
def cargar():
    # Ladrillo: variable texto que dice qué botón se presionó ("ver" o "descargar")
    accion = request.form.get("accion", "ver")
    # Ladrillo: los dos archivos que llegan del formulario
    archivo_maestro = request.files.get("maestro")
    archivo_evaluaciones = request.files.get("evaluaciones")
    error = None  # Ladrillo: variable texto para el aviso rojo

    try:
        # Ladrillo: CONDICIONAL que revisa que los dos archivos hayan llegado
        if (not archivo_maestro or not archivo_maestro.filename
                or not archivo_evaluaciones or not archivo_evaluaciones.filename):
            error = "Debe seleccionar los dos archivos: el Maestro de estudiantes y el Registro de evaluaciones."
        # Ladrillo: CONDICIONAL que revisa la extensión
        elif (not archivo_maestro.filename.lower().endswith(".xlsx")
              or not archivo_evaluaciones.filename.lower().endswith(".xlsx")):
            error = "Los dos archivos deben ser Excel con extensión .xlsx."
        else:
            # Ladrillo: leer el contenido a bytes (en memoria, nunca a disco)
            contenido_maestro = archivo_maestro.read()
            contenido_evaluaciones = archivo_evaluaciones.read()
            limite = 4 * 1024 * 1024  # Ladrillo: 4 MB en bytes
            # Ladrillo: CONDICIONAL de tamaño por archivo
            if len(contenido_maestro) > limite or len(contenido_evaluaciones) > limite:
                error = "Cada archivo debe pesar menos de 4 MB. Suba archivos más livianos."
            else:
                # Ladrillo: BytesIO = la "memoria USB imaginaria" dentro de la RAM.
                # openpyxl lee de ahí igual que de un archivo del disco, pero nada se guarda.
                memoria_maestro = io.BytesIO(contenido_maestro)
                memoria_evaluaciones = io.BytesIO(contenido_evaluaciones)
                resultados, contadores, inconsistencias, programas = cargar_datos(
                    memoria_maestro, memoria_evaluaciones)

                # Ladrillo: CONDICIONAL según el botón presionado
                if accion == "descargar":
                    buffer = construir_excel(resultados)
                    return send_file(
                        buffer,
                        as_attachment=True,
                        download_name="resultados_certificados.xlsx",
                        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                # Ladrillo: origen = nombres de los archivos subidos
                origen = f"Archivos subidos: {archivo_maestro.filename} y {archivo_evaluaciones.filename}"
                return pintar_tablero(resultados, contadores, inconsistencias, programas,
                                      origen=origen, aviso_subida=True)
    # Ladrillo: si algo falla al leer el Excel, no se deja la página en blanco
    except Exception:
        app.logger.exception("Error al procesar los Excel subidos")
        error = ("No se pudieron leer los Excel. Verifique que sean .xlsx válidos y que tengan "
                 "las hojas 'Estudiantes' y 'Evaluaciones'.")

    # Ladrillo: si hubo error, se muestra el tablero de fábrica con el aviso rojo
    resultados, contadores, inconsistencias, programas = cargar_datos()
    return pintar_tablero(resultados, contadores, inconsistencias, programas, error=error)


# Ladrillo: descarga de los resultados con los datos de fábrica (solo admin)
@app.route("/descargar")
@requiere_login
@solo_admin
def descargar():
    resultados, _, _, _ = cargar_datos()
    buffer = construir_excel(resultados)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="resultados_certificados.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# Ladrillo: si la petición pesa más de 4 MB, Flask corta antes y aquí se avisa con claridad
@app.errorhandler(413)
def archivo_demasiado_grande(_error):
    # Ladrillo: CONDICIONAL, sin sesión no se muestra nada
    if "correo" not in session:
        return redirect(url_for("login"))
    resultados, contadores, inconsistencias, programas = cargar_datos()
    return pintar_tablero(
        resultados, contadores, inconsistencias, programas,
        error="La carga supera el límite de 4 MB en total (Vercel rechaza más de 4.5 MB). "
              "Suba archivos más livianos o procéselos por separado."
    ), 413


# Ladrillo: ruta de detalle por estudiante (solo muestra, no calcula nada nuevo), protegida
@app.route("/estudiante/<identificacion>")
@requiere_login
def estudiante(identificacion):
    resultados, _, _, _ = cargar_datos()
    # Ladrillo: BUCLE para buscar a la persona por su identificación (texto)
    for r in resultados:
        # Ladrillo: CONDICIONAL que compara textos
        if r["identificacion"] == str(identificacion).strip():
            return render_template("estudiante.html", r=r)
    return render_template("estudiante.html", r=None)


# Ladrillo: arranque de la app (solo se ejecuta al abrir este archivo directo)
if __name__ == "__main__":
    app.run(debug=True)
